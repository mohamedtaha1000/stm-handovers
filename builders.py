#!/usr/bin/env python3
"""
builders.py
===========
One function per SHAPE of handover document - not per document type.
Ten types, three shapes: a laptop handover, a laptop replacement (two
spec tables, old and new), and everything else, which is the same
employee block plus one spec table whatever the accessory is.

Each takes (template_path, output_path, data) and writes the filled
file. They are the only place that knows which FILL_ token belongs to
which field; the mechanics of editing the XML live in ooxml.py, and what
fields exist at all lives in templates.py.
"""

from collections.abc import Callable
from pathlib import Path
from typing import Any

import docx

from ooxml import (
    date_parts,
    drop_empty_columns,
    fill_placeholders,
    find_spec_tables,
    fit_to_page,
    iter_paragraphs,
    level_spec_rows,
    set_spec_cell,
    spec_table,
    tidy_document,
)
from templates import DEFAULT_COMPANY

# ----------------------------------------------------------------------
# Laptop handover
# ----------------------------------------------------------------------

def fill_laptop_handover(template_path: Path, output_path: Path, data: dict[str, Any]) -> Path:
    """Laptop handover.

    Template refreshed again 2026-09 (v5): the source .docx is now
    placeholder-driven (FILL_NAME, FILL_SERIAL, ...), so this function no
    longer walks tables by row/column/run index. That indexing is exactly
    what broke when the document was last re-saved from Word - the labels
    and values were re-split into different runs and the hand-counted
    offsets wrote into the wrong places (or raised IndexError).

    Color / Hard / RAM are not tokens in this template - they carry the
    common defaults as literal text - so they're written by column name
    afterwards from whatever the form actually submitted."""
    d = docx.Document(str(template_path))
    date_str, day_name, month_name = date_parts(data["date_obj"])

    email = (data.get("email") or "").strip()
    company = data.get("company") or DEFAULT_COMPANY

    values = {
        "FILL_DATE": date_str,
        "FILL_MONTH": month_name,
        "FILL_DAY": day_name,
        "FILL_DAYNAME": day_name,
        "FILL_NAME": data["name"],
        "FILL_ROLE": data["role"],
        "FILL_DEPARTMENT": data["department"],
        "FILL_MOBILE": data["mobile"],
        "FILL_EMAIL": email,
        "FILL_CODE": data["code"],
        "FILL_GOVID": data["govid"],
        "FILL_COMPANY": company,
        "FILL_MODEL": data["model"],
        "FILL_SERIAL": data["serial"],
        "FILL_CPU": data["cpu"],
        # Present in some revisions of the template, literal defaults in
        # others - mapped either way so both revisions fill correctly.
        "FILL_COLOR": data.get("color", ""),
        "FILL_HARD": data.get("storage", ""),
        "FILL_STORAGE": data.get("storage", ""),
        "FILL_RAM": data.get("ram", ""),
    }
    specs = find_spec_tables(d)
    tidy_document(d, values)
    fill_placeholders(d, values)
    level_spec_rows(specs)
    for spec in specs:
        drop_empty_columns(spec)
    fit_to_page(d, specs)

    specs = spec_table(d, "serial", "cpu")
    set_spec_cell(specs, "Color", data.get("color", ""))
    set_spec_cell(specs, "Hard", data.get("storage", ""))
    set_spec_cell(specs, "RAM", data.get("ram", ""))

    # The accessories note under the signature block lists the laptop
    # charger, bag and mouse. Older source documents hardcoded one real
    # mouse brand/model there, which would print on every document
    # regardless of what was issued - strip it if it comes back.
    for p in iter_paragraphs(d):
        for run in p.runs:
            if "Logitech" in run.text or "M171" in run.text:
                run.text = ""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(output_path))
    return output_path


# ----------------------------------------------------------------------
# Laptop replacement
# ----------------------------------------------------------------------

def fill_laptop_replacement(template_path: Path, output_path: Path, data: dict[str, Any]) -> Path:
    """Laptop replacement (old device returned + new device issued).

    Same story as the handover above - the refreshed source document is
    placeholder-driven, so the old index-based table walking is gone. It
    also quietly fixes a swap that the index version had wrong: in this
    revision the FIRST spec table is the OLD device and the second is the
    NEW one, the opposite of the previous file's order, which is how old
    and new specs ended up printed on each other's rows."""
    d = docx.Document(str(template_path))
    date_str, day_name, month_name = date_parts(data["date_obj"])

    email = (data.get("email") or "").strip()
    company = data.get("company") or DEFAULT_COMPANY

    values = {
        "FILL_DATE": date_str,
        "FILL_MONTH": month_name,
        "FILL_DAYNAME": day_name,
        "FILL_DAY": day_name,
        "FILL_NAME": data["name"],
        "FILL_ROLE": data["role"],
        "FILL_DEPARTMENT": data["department"],
        "FILL_MOBILE": data["mobile"],
        "FILL_EMAIL": email,
        "FILL_CODE": data["code"],
        "FILL_GOVID": data["govid"],
        "FILL_COMPANY": company,
        "FILL_OLD_MODEL": data["old_model"],
        "FILL_OLD_SERIAL": data["old_serial"],
        "FILL_OLD_COLOR": data["old_color"],
        "FILL_OLD_STORAGE": data["old_storage"],
        "FILL_OLD_HARD": data["old_storage"],
        "FILL_OLD_RAM": data["old_ram"],
        "FILL_OLD_CPU": data["old_cpu"],
        "FILL_NEW_MODEL": data["new_model"],
        "FILL_NEW_SERIAL": data["new_serial"],
        "FILL_NEW_COLOR": data["new_color"],
        "FILL_NEW_STORAGE": data["new_storage"],
        "FILL_NEW_HARD": data["new_storage"],
        "FILL_NEW_RAM": data["new_ram"],
        "FILL_NEW_CPU": data["new_cpu"],
    }
    specs = find_spec_tables(d)
    tidy_document(d, values)
    fill_placeholders(d, values)
    level_spec_rows(specs)
    for spec in specs:
        drop_empty_columns(spec)
    fit_to_page(d, specs)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(output_path))
    return output_path


# ----------------------------------------------------------------------
# The eight accessory documents (mouse, keyboard & mouse, screen, router,
# headset, printer, flash drive, hard disk)
#
# These were converted to the same placeholder form as the laptop
# templates, and they all share one six-table skeleton - title, date
# header, employee details, device spec, declaration, signatures. That
# makes one fill function enough for all eight: every difference between
# them (a router has a SIM number, a flash drive has a capacity, a mouse
# has neither) is expressed by which placeholders that document contains,
# so the mapping below can simply offer all of them and let each template
# take the ones it has. Adding a ninth accessory document later needs a
# registry entry and a .docx, and no new code at all.
# ----------------------------------------------------------------------

def fill_accessory_handover(template_path: Path, output_path: Path, data: dict[str, Any]) -> Path:
    """Fill any of the accessory handover documents."""
    d = docx.Document(str(template_path))
    date_str, day_name, month_name = date_parts(data["date_obj"])

    values: dict[str, Any] = {
        "FILL_DATE": date_str,
        "FILL_MONTH": month_name,
        "FILL_DAY": day_name,
        "FILL_DAYNAME": day_name,
        "FILL_NAME": data["name"],
        "FILL_ROLE": data["role"],
        "FILL_DEPARTMENT": data["department"],
        "FILL_MOBILE": data["mobile"],
        "FILL_EMAIL": (data.get("email") or "").strip(),
        "FILL_CODE": data["code"],
        "FILL_GOVID": data["govid"],
        "FILL_COMPANY": data.get("company") or DEFAULT_COMPANY,
    }
    # Device fields, offered for every template; each document fills in
    # only the ones its own spec table actually asks for.
    for key, token in (
        ("model", "FILL_MODEL"), ("serial", "FILL_SERIAL"),
        ("brand", "FILL_BRAND"), ("color", "FILL_COLOR"),
        ("capacity", "FILL_CAPACITY"), ("sim_number", "FILL_SIM_NUMBER"),
        ("storage", "FILL_STORAGE"), ("ram", "FILL_RAM"), ("cpu", "FILL_CPU"),
    ):
        if key in data:
            values[token] = data[key]

    specs = find_spec_tables(d)
    tidy_document(d, values)
    fill_placeholders(d, values)
    level_spec_rows(specs)
    for spec in specs:
        drop_empty_columns(spec)
    fit_to_page(d, specs)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(output_path))
    return output_path


# Which shape each document type is filled by. Kept here rather than in
# the registry so templates.py stays what it says it is: data, with no
# imports and nothing to execute.
FILL_FUNCTIONS = {
    "laptop_handover": fill_laptop_handover,
    "laptop_replacement": fill_laptop_replacement,
}


def fill_function(template_id: str) -> Callable[[Path, Path, dict[str, Any]], Path]:
    """How to fill this document type. Everything that is not a laptop
    is the accessory shape."""
    return FILL_FUNCTIONS.get(template_id, fill_accessory_handover)
